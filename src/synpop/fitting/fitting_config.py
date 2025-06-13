# pylint: disable=too-many-lines
"""Utilities for setting up the fitting config files."""
import logging
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Dict
from typing import List
from typing import Optional
from typing import Union

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles.fonts import Font

from synpop.validation import get_categories


# Logger settings set in __init__.py
logger = logging.getLogger(__name__)


def write_fitting_goal(
    excel_path: str,
    expected_counts: pd.DataFrame,
    target_variable: str,
    fitting_segments: List[str],
    proba_weights: str = "uniform",
    overwrite_file: bool = False,
    sheet_name: Optional[str] = None,
    fit_as_delta: bool = False,
) -> None:
    """Write given fitting goal to the fitting excel file.

    This is a helper function. The same can also be achieved by hand.
    """
    if not sheet_name:
        sheet_name = target_variable
    if target_variable in ("level_of_employment", "is_employed"):
        raise ValueError('Fitting employment only supported for "current_job_rank".')

    mode = "a" if Path(excel_path).exists() and not overwrite_file else "w"
    if isinstance(fitting_segments, list):
        fitting_segments = ",".join(fitting_segments)  # type: ignore

    if isinstance(proba_weights, list):
        proba_weights = ",".join(proba_weights)

    with pd.ExcelWriter(  # pylint: disable=abstract-class-instantiated
        excel_path, engine="openpyxl", mode=mode
    ) as writer:  # pylint: disable=abstract-class-instantiated
        # write control totals (if sheet already exists, it shall be overwritten)
        try:
            writer.book.remove(writer.book[sheet_name])  # pylint: disable=no-member
        except IOError:
            pass
        finally:
            expected_counts.to_excel(writer, sheet_name=sheet_name)

        # insert new rows. If there are merged cells (MultiIndex), shift them first
        sheet = writer.sheets[sheet_name]
        merged_cells_range = sheet.merged_cells.ranges
        for merged_cell in merged_cells_range:
            merged_cell.shift(0, 4)
        sheet.insert_rows(idx=0, amount=4)

        # fill config and format
        sheet["A1"] = "target_variable"
        sheet["A2"] = "fitting_segments"
        sheet["A3"] = "probability_weights"
        sheet["A4"] = "fitting_type"
        sheet["B1"] = target_variable
        sheet["B2"] = fitting_segments
        sheet["B3"] = proba_weights
        sheet["B4"] = "delta" if fit_as_delta else "absolute"
        sheet["A1"].font = Font(bold=True)
        sheet["A2"].font = Font(bold=True)
        sheet["A3"].font = Font(bold=True)
        sheet.column_dimensions["A"].width = 20

        # finish and close
        writer.close()  # type: ignore


@dataclass
class VariableFittingConfig:
    """Configuration for fitting a single target variable."""

    target_variable: str
    fitting_segments: List[str]
    probability_weights: List[str]
    fitting_type: str

    @property
    def fit_as_delta(self) -> bool:
        return self.fitting_type == "delta"


class FittingConfig:
    """Configuration and counts for fitting target variables."""

    config: VariableFittingConfig
    expected_counts: pd.DataFrame
    _new_format: bool = True

    person_categories: Dict[str, List[str]] = get_categories("persons")
    person_categories.update(
        {
            "pop_total": ["pop_total"],
            "jobs_endo": ["jobs_endo"],
            "jobs_exo": ["jobs_exo"],
        }
    )

    @staticmethod
    def from_excel_sheet(excel_path, sheetname):
        """Read a fitting configuration from an excel sheet."""
        fitting_config = FittingConfig()
        fitting_config.config = fitting_config.parse_fitting_config(
            excel_path, sheetname
        )
        fitting_config.expected_counts = fitting_config.parse_fitting_goals(
            excel_path, sheetname
        )
        return fitting_config

    def parse_fitting_config(
        self, excel_path: Union[Path, str], sheetname: str
    ) -> VariableFittingConfig:
        """Parse the fitting configuration from an excel sheet."""
        config = pd.read_excel(
            excel_path,
            sheet_name=sheetname,
            engine="openpyxl",
            usecols=[0, 1],
            header=None,
            nrows=4,
            names=["config_key", "config_value"],
            index_col=0,
        )["config_value"].to_dict()
        config["fitting_segments"] = config["fitting_segments"].split(",")
        config["probability_weights"] = config["probability_weights"].split(",")
        if "uniform" in config["probability_weights"][0]:
            config["probability_weights"] = None
        elif config["target_variable"] in ("pop_total", "jobs_endo", "jobs_exo"):
            raise NotImplementedError(
                f"Weighted fitting for {config['target_variable']} not supported."
            )
        if "fitting_type" not in config:
            config["fitting_type"] = "absolute"
            self._new_format = False

        return VariableFittingConfig(**config)

    def parse_fitting_goals(
        self, excel_path: Union[Path, str], sheetname: str
    ) -> pd.DataFrame:
        """Parse the fitting goals from an excel sheet."""
        expected_counts = pd.read_excel(
            excel_path,
            sheet_name=sheetname,
            engine="openpyxl",
            skiprows=4 + int(self._new_format),
            na_values=["NaN", ""],
            keep_default_na=False,
            index_col=list(range(len(self.config.fitting_segments))),
        )
        usecols = [
            c
            for c in expected_counts.columns
            if c
            in self.config.fitting_segments
            + self.person_categories[self.config.target_variable]
        ]
        expected_counts = expected_counts[usecols].dropna().astype(int)
        return expected_counts


def parse_fitting_goals(
    excel_path: Union[Path, str],
) -> Dict[str, List[FittingConfig]]:
    """Parse the fitting goals from an excel file."""
    sheetnames = load_workbook(excel_path).sheetnames
    configs = defaultdict(list)
    for sheet in sheetnames:
        fitting_config = FittingConfig.from_excel_sheet(excel_path, sheet)
        configs[fitting_config.config.target_variable].append(fitting_config)

    return configs
